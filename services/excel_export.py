"""
Export Excel du portefeuille (V0.7). Utilise le même modèle de colonnes que l'import
(services/import_export.EXCEL_COLUMNS) pour garantir un aller-retour cohérent.
"""
import io
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from services.import_export import EXCEL_COLUMNS


def _get_path(d, path):
    v = d
    for key in path:
        v = v.get(key) if isinstance(v, dict) else None
    return v


def build_export_workbook(services, kpis, organisation_name, application_name):
    wb = Workbook()

    # ---------- Feuille Synthèse ----------
    ws_synthese = wb.active
    ws_synthese.title = "Synthèse"

    title_font = Font(name="Arial", size=16, bold=True, color="0F4C81")
    subtitle_font = Font(name="Arial", size=11, color="4B5568")
    header_font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="0F4C81", end_color="0F4C81", fill_type="solid")
    label_font = Font(name="Arial", size=10, bold=True)
    value_font = Font(name="Arial", size=10)

    ws_synthese["B2"] = application_name
    ws_synthese["B2"].font = title_font
    ws_synthese["B3"] = f"Portefeuille des services SI — {organisation_name}"
    ws_synthese["B3"].font = subtitle_font
    ws_synthese["B4"] = f"Généré le {datetime.now().strftime('%d/%m/%Y à %H:%M')}"
    ws_synthese["B4"].font = subtitle_font

    kpi_rows = [
        ("Total services", kpis.get("total", 0)),
        ("En catalogue (production)", kpis.get("par_statut_portfolio", {}).get("Catalogue", 0)),
        ("En pipeline", kpis.get("par_statut_portfolio", {}).get("Pipeline", 0)),
        ("Criticité critique", kpis.get("par_criticite", {}).get("Critique", 0)),
        ("Couverture gouvernance (%)", kpis.get("couverture_gouvernance_pct", 0)),
        ("Couverture continuité RTO/RPO (%)", kpis.get("couverture_continuite_pct", 0)),
        ("Dépendance prestataire (%)", kpis.get("taux_dependance_prestataire_pct", 0)),
        ("Complétude moyenne (%)", kpis.get("completude_moyenne_pct", 0)),
    ]

    row = 6
    ws_synthese[f"B{row}"] = "Indicateur"
    ws_synthese[f"C{row}"] = "Valeur"
    for col in ("B", "C"):
        ws_synthese[f"{col}{row}"].font = header_font
        ws_synthese[f"{col}{row}"].fill = header_fill
    row += 1
    for label, value in kpi_rows:
        ws_synthese[f"B{row}"] = label
        ws_synthese[f"B{row}"].font = label_font
        ws_synthese[f"C{row}"] = value
        ws_synthese[f"C{row}"].font = value_font
        row += 1

    ws_synthese.column_dimensions["A"].width = 3
    ws_synthese.column_dimensions["B"].width = 34
    ws_synthese.column_dimensions["C"].width = 16

    # ---------- Feuille Inventaire ----------
    ws_inv = wb.create_sheet("Inventaire")
    headers = [name for name, _ in EXCEL_COLUMNS]

    for col_idx, name in enumerate(headers, start=1):
        cell = ws_inv.cell(row=1, column=col_idx, value=name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    data_font = Font(name="Arial", size=10)
    for row_idx, svc in enumerate(services, start=2):
        for col_idx, (name, path) in enumerate(EXCEL_COLUMNS, start=1):
            value = svc.get("id") if path == ("id",) else _get_path(svc, path)
            cell = ws_inv.cell(row=row_idx, column=col_idx, value=value)
            cell.font = data_font

    for col_idx, name in enumerate(headers, start=1):
        ws_inv.column_dimensions[get_column_letter(col_idx)].width = max(14, len(name) + 2)

    ws_inv.freeze_panes = "A2"
    if services:
        ws_inv.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(services) + 1}"

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer
