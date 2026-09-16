"""Tests du module import/export (V0.7) : parsing JSON/Excel, preview, stratégies, export xlsx."""
import os
import sys
import io
import json

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import pytest
import openpyxl

from services.import_export import parse_json, parse_excel, build_preview, apply_import, EXCEL_COLUMNS, ImportError_
from services.excel_export import build_export_workbook
from services.service_manager import ServiceManager

FIXTURE_FILE = os.path.join(os.path.dirname(os.path.abspath(__file__)), "fixtures", "sample_services.json")


def load_fixture_services():
    with open(FIXTURE_FILE, encoding="utf-8") as f:
        return json.load(f)["services"]


def make_json_bytes(services):
    return json.dumps({"metadata": {}, "services": services}).encode("utf-8")


def make_excel_bytes(rows):
    """rows: liste de dicts {colonne: valeur} selon EXCEL_COLUMNS."""
    wb = openpyxl.Workbook()
    ws = wb.active
    headers = [name for name, _ in EXCEL_COLUMNS]
    ws.append(headers)
    for row in rows:
        ws.append([row.get(h) for h in headers])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------- parse_json ----------

def test_parse_json_valid():
    services = load_fixture_services()
    parsed, errors = parse_json(make_json_bytes(services))
    assert len(parsed) == 5
    assert errors == []


def test_parse_json_missing_nom_reported_as_error():
    services = [{"id": "APP-X01", "identification": {}}]
    parsed, errors = parse_json(make_json_bytes(services))
    assert len(parsed) == 0
    assert len(errors) == 1


def test_parse_json_invalid_structure_raises():
    with pytest.raises(ImportError_):
        parse_json(b'{"pas_de_cle_services": true}')


def test_parse_json_malformed_raises():
    with pytest.raises(ImportError_):
        parse_json(b"{ceci n'est pas du json")


# ---------- parse_excel ----------

def test_parse_excel_valid_rows():
    rows = [
        {"ID": "APP-X01", "Nom": "Service Excel 1", "Criticite": "Critique"},
        {"ID": "APP-X02", "Nom": "Service Excel 2", "Hebergement": "Cloud"},
    ]
    parsed, errors = parse_excel(make_excel_bytes(rows))
    assert len(parsed) == 2
    assert errors == []
    assert parsed[0]["identification"]["nom"] == "Service Excel 1"
    assert parsed[0]["meta"]["criticite"] == "Critique"


def test_parse_excel_missing_nom_reported():
    rows = [{"ID": "APP-X01", "Nom": None}]
    parsed, errors = parse_excel(make_excel_bytes(rows))
    assert len(parsed) == 0
    assert len(errors) == 1


def test_parse_excel_without_nom_column_raises():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["ID", "Criticite"])
    ws.append(["APP-X01", "Critique"])
    buf = io.BytesIO()
    wb.save(buf)
    with pytest.raises(ImportError_):
        parse_excel(buf.getvalue())


# ---------- build_preview ----------

def test_preview_categorizes_new_and_existing():
    existing = load_fixture_services()
    imported = [
        {**ServiceManager.empty_skeleton(), "id": "APP-T01", "identification": {**ServiceManager.empty_skeleton()["identification"], "nom": "Portail Démo Modifié"}, "meta": {"statut": None, "criticite": None, "completude_pct": 0}},
    ]
    preview = build_preview(existing, imported, [])
    assert preview["resume"]["modifies"] == 1
    assert preview["resume"]["nouveaux"] == 0


def test_preview_detects_new_service():
    existing = load_fixture_services()
    new_svc = ServiceManager.empty_skeleton()
    new_svc["id"] = "APP-NEW"
    new_svc["identification"]["nom"] = "Tout Nouveau Service"
    preview = build_preview(existing, [new_svc], [])
    assert preview["resume"]["nouveaux"] == 1


def test_preview_detects_duplicate_within_file():
    svc1 = ServiceManager.empty_skeleton()
    svc1["id"] = "APP-DUP"
    svc1["identification"]["nom"] = "Premier"
    svc2 = ServiceManager.empty_skeleton()
    svc2["id"] = "APP-DUP"
    svc2["identification"]["nom"] = "Deuxième"
    preview = build_preview([], [svc1, svc2], [])
    assert preview["resume"]["doublons_fichier"] == 1


# ---------- apply_import ----------

def test_apply_import_remplacer():
    existing = load_fixture_services()
    new_svc = ServiceManager.empty_skeleton()
    new_svc["id"] = "APP-ONLY"
    new_svc["identification"]["nom"] = "Seul Service"
    result = apply_import(existing, [new_svc], "remplacer")
    assert len(result) == 1
    assert result[0]["id"] == "APP-ONLY"


def test_apply_import_ajouter_skips_existing_ids():
    existing = load_fixture_services()
    dup = ServiceManager.empty_skeleton()
    dup["id"] = "APP-T01"
    dup["identification"]["nom"] = "Ne doit pas remplacer"
    new = ServiceManager.empty_skeleton()
    new["id"] = "APP-NEW2"
    new["identification"]["nom"] = "Nouveau via ajout"

    result = apply_import(existing, [dup, new], "ajouter")
    assert len(result) == len(existing) + 1
    kept = next(s for s in result if s["id"] == "APP-T01")
    assert kept["identification"]["nom"] == "Portail Démo"  # non écrasé


def test_apply_import_maj_updates_existing_and_adds_new():
    existing = load_fixture_services()
    upd = ServiceManager.empty_skeleton()
    upd["id"] = "APP-T01"
    upd["identification"]["nom"] = "Portail Démo Mis À Jour"
    new = ServiceManager.empty_skeleton()
    new["id"] = "APP-NEW3"
    new["identification"]["nom"] = "Ajouté via maj"

    result = apply_import(existing, [upd, new], "maj")
    assert len(result) == len(existing) + 1
    updated = next(s for s in result if s["id"] == "APP-T01")
    assert updated["identification"]["nom"] == "Portail Démo Mis À Jour"


def test_apply_import_invalid_strategy_raises():
    with pytest.raises(ImportError_):
        apply_import([], [], "inconnue")


# ---------- export Excel ----------

def test_export_workbook_structure():
    services = load_fixture_services()
    kpis = {"total": len(services), "par_statut_portfolio": {}, "par_criticite": {},
            "couverture_gouvernance_pct": 60, "couverture_continuite_pct": 20,
            "taux_dependance_prestataire_pct": 40, "completude_moyenne_pct": 41}

    buffer = build_export_workbook(services, kpis, "Organisation Test", "MIRSAAD")
    wb = openpyxl.load_workbook(buffer)

    assert "Synthèse" in wb.sheetnames
    assert "Inventaire" in wb.sheetnames

    ws_inv = wb["Inventaire"]
    assert ws_inv.cell(row=1, column=1).value == "ID"
    assert ws_inv.max_row == len(services) + 1  # + ligne d'en-tête


def test_export_import_roundtrip_preserves_key_fields():
    services = load_fixture_services()
    kpis = {"total": len(services)}
    buffer = build_export_workbook(services, kpis, "Org", "MIRSAAD")

    parsed, errors = parse_excel(buffer.getvalue())
    assert errors == []
    assert len(parsed) == len(services)
    noms_export = {s["identification"]["nom"] for s in services}
    noms_reimport = {s["identification"]["nom"] for s in parsed}
    assert noms_export == noms_reimport
